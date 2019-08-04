import pandas as pd
import openpyxl # extract hyperlinks from an excel file

def get_hyperlinks_col(file):
    wb = openpyxl.load_workbook(file)
    ws = wb.get_sheet_by_name('Результаты поиска')
    hyperlinks = []
    for i in range(ws.max_row - 3):
        url = ws.cell(row = i + 3, column = 27).value.split(',')[0].replace('=HYPERLINK(', '').replace('"', '')
        hyperlinks.append(url)
    return hyperlinks

def file_transform(file):
    hyperlinks = get_hyperlinks_col(file)
    df = pd.read_excel(file, 'Результаты поиска', skiprows=2, header=None)
    df.drop(df.columns[[0,5,11,12,13,16,17,18,19,20,22,23,24,25,27,28]], axis=1, inplace=True) # drop non-informative cols
    df.columns = ['client_name', 'client_inn', 'contract_sum', 'contract_id', 'okpd', 'contract_name', 'region', 'city',
              'start_date', 'supplier_name', 'supplier_inn', 'auction_type', 'legislation']
    df['hyperlinks'] = hyperlinks
    df.client_inn = df.client_inn.astype('str')
    df.supplier_inn = df.supplier_inn.astype('str')
    return df
